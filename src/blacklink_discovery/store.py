import csv
import hashlib
import json
import os
import shutil
import sqlite3
import tempfile
import uuid
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from .extract import extract

SCHEMA = """
CREATE TABLE IF NOT EXISTS runs(id INTEGER PRIMARY KEY, started TEXT DEFAULT CURRENT_TIMESTAMP, label TEXT);
CREATE TABLE IF NOT EXISTS resources(id INTEGER PRIMARY KEY, run_id INTEGER, site TEXT, url TEXT,
 kind TEXT, status INTEGER, evidence_path TEXT, error TEXT);
CREATE TABLE IF NOT EXISTS text_views(resource_id INTEGER PRIMARY KEY,path TEXT);
CREATE TABLE IF NOT EXISTS edges(run_id INTEGER,site TEXT,parent TEXT,child TEXT,relation TEXT,
 parent_resource_id INTEGER DEFAULT 0, start INTEGER, end INTEGER,
 UNIQUE(run_id,site,parent,child,relation,parent_resource_id,start,end));
CREATE TABLE IF NOT EXISTS findings(resource_id INTEGER,host TEXT,registered TEXT,host_type TEXT,
 start INTEGER,end INTEGER,line INTEGER,col INTEGER,source TEXT,method TEXT,evidence TEXT,
 UNIQUE(resource_id,host,start,end,source));
CREATE TABLE IF NOT EXISTS queue(run_id INTEGER,site TEXT,url TEXT,depth INTEGER,asset INTEGER,
 state TEXT DEFAULT 'pending',error TEXT, UNIQUE(run_id,site,url));
CREATE TABLE IF NOT EXISTS run_configs(run_id INTEGER PRIMARY KEY,config TEXT);
CREATE TABLE IF NOT EXISTS run_events(id INTEGER PRIMARY KEY,run_id INTEGER,
 time TEXT DEFAULT CURRENT_TIMESTAMP,event TEXT,detail TEXT);
CREATE TABLE IF NOT EXISTS attempts(run_id INTEGER,site TEXT,url TEXT,count INTEGER DEFAULT 0,
 UNIQUE(run_id,site,url));
CREATE TABLE IF NOT EXISTS request_observations(id INTEGER PRIMARY KEY,run_id INTEGER,site TEXT,
 page TEXT,url TEXT,method TEXT,resource_type TEXT,state TEXT,status INTEGER,error TEXT,
 resource_id INTEGER);
CREATE TABLE IF NOT EXISTS browser_fallbacks(run_id INTEGER,site TEXT,url TEXT,browser_run_id INTEGER,
 state TEXT,UNIQUE(run_id,site,url));
CREATE TABLE IF NOT EXISTS expansions(resource_id INTEGER PRIMARY KEY,depth INTEGER);
CREATE TABLE IF NOT EXISTS browser_depths(run_id INTEGER,url TEXT,depth INTEGER,
 UNIQUE(run_id,url));
CREATE TABLE IF NOT EXISTS static_requests(id INTEGER PRIMARY KEY,run_id INTEGER,site TEXT,url TEXT,
 fingerprint TEXT,status INTEGER,resource_id INTEGER,error TEXT,time TEXT DEFAULT CURRENT_TIMESTAMP);
CREATE INDEX IF NOT EXISTS static_fingerprint_attempts ON static_requests(run_id,site,url,fingerprint,status);
"""

EXCEL_ROWS = 1048576


def _atomic_blob(path: Path, content: bytes):
    if path.is_file():
        if hashlib.sha256(path.read_bytes()).digest() != hashlib.sha256(content).digest():
            raise ValueError("existing evidence blob failed integrity check")
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(dir=path.parent)
    try:
        with os.fdopen(fd, "wb") as stream:
            stream.write(content)
            stream.flush()
            os.fsync(stream.fileno())
        try:
            os.link(tmp, path)
        except FileExistsError:
            pass
    finally:
        Path(tmp).unlink(missing_ok=True)


class Store:
    def __init__(self, directory: str, read_only: bool = False):
        self.directory = Path(directory)
        self.read_only = read_only
        if read_only:
            database = self.directory / "discovery.sqlite"
            if not database.is_file():
                raise ValueError("results database does not exist; check --output")
            self.db = sqlite3.connect(database.resolve().as_uri() + "?mode=ro", uri=True)
        else:
            self.directory.mkdir(parents=True, exist_ok=True)
            self.db = sqlite3.connect(self.directory / "discovery.sqlite")
        self.db.row_factory = sqlite3.Row
        if not read_only:
            self.db.executescript(SCHEMA)
            self._migrate()

    def _migrate(self):
        self.db.execute("BEGIN IMMEDIATE")
        indexes = self.db.execute("PRAGMA index_list(resources)").fetchall()
        unique_url = False
        for index in indexes:
            if index[2] and [r[2] for r in self.db.execute(f'PRAGMA index_info("{index[1]}")')] == [
                "run_id",
                "site",
                "url",
            ]:
                unique_url = True
        if unique_url:
            self.db.execute("ALTER TABLE resources RENAME TO resources_legacy")
            self.db.execute("""CREATE TABLE resources(id INTEGER PRIMARY KEY,run_id INTEGER,site TEXT,url TEXT,
                kind TEXT,status INTEGER,evidence_path TEXT,error TEXT)""")
            self.db.execute("INSERT INTO resources SELECT * FROM resources_legacy")
            self.db.execute("DROP TABLE resources_legacy")
        columns = {r[1] for r in self.db.execute("PRAGMA table_info(resources)")}
        for name, definition in {
            "collector": "TEXT DEFAULT 'legacy'",
            "encoding": "TEXT DEFAULT ''",
            "parent_resource_id": "INTEGER DEFAULT 0",
        }.items():
            if name not in columns:
                self.db.execute(f"ALTER TABLE resources ADD COLUMN {name} {definition}")
        columns = {r[1] for r in self.db.execute("PRAGMA table_info(edges)")}
        if "parent_resource_id" not in columns:
            self.db.execute("ALTER TABLE edges ADD COLUMN parent_resource_id INTEGER DEFAULT 0")
        # Rebuild the old URL-only edge uniqueness without deleting legacy rows.
        definition = self.db.execute("SELECT sql FROM sqlite_master WHERE name='edges'").fetchone()[
            0
        ]
        # Each source offset is a separate occurrence. Rebuilding adds the columns
        # without collapsing legacy rows that had no recorded position.
        if "start" not in definition.replace(" ", "").replace("\n", ""):
            self.db.execute("ALTER TABLE edges RENAME TO edges_legacy")
            self.db.execute(
                "CREATE TABLE edges(run_id INTEGER,site TEXT,parent TEXT,child TEXT,relation TEXT,parent_resource_id INTEGER DEFAULT 0,start INTEGER,end INTEGER,UNIQUE(run_id,site,parent,child,relation,parent_resource_id,start,end))"
            )
            legacy = {r[1] for r in self.db.execute("PRAGMA table_info(edges_legacy)")}
            start = "start" if "start" in legacy else "NULL"
            end = "end" if "end" in legacy else "NULL"
            self.db.execute(
                "INSERT INTO edges SELECT run_id,site,parent,child,relation,COALESCE(parent_resource_id,0),"
                f"{start},{end} FROM edges_legacy"
            )
            self.db.execute("DROP TABLE edges_legacy")
        self.db.execute(
            "CREATE INDEX IF NOT EXISTS resource_lookup ON resources(run_id,site,url,evidence_path)"
        )
        self.db.execute(
            "CREATE INDEX IF NOT EXISTS queue_pending_schedule ON queue(run_id,state,depth)"
        )
        self.db.commit()

    def new_run(self, label: str) -> int:
        cursor = self.db.execute("INSERT INTO runs(label) VALUES(?)", (label,))
        self.db.commit()
        return cursor.lastrowid

    def edge(
        self,
        run: int,
        site: str,
        parent: str,
        child: str,
        relation: str,
        parent_resource_id=0,
        start=None,
        end=None,
    ):
        self.db.execute(
            "INSERT OR IGNORE INTO edges(run_id,site,parent,child,relation,parent_resource_id,start,end) VALUES(?,?,?,?,?,?,?,?)",
            (run, site, parent, child, relation, parent_resource_id, start, end),
        )
        self.db.commit()

    def save(
        self,
        run: int,
        site: str,
        url: str,
        body: bytes,
        text: str,
        kind: str,
        status: int = 200,
        error: str = "",
        *,
        parent_resource_id=0,
        collector="static",
        encoding="",
        mentions=None,
    ):
        mentions = extract(text, kind) if mentions is None else mentions
        digest = hashlib.sha256(body).hexdigest()
        decoded = text.encode("utf-8")
        text_digest = hashlib.sha256(decoded).hexdigest()
        body_path = self.directory / "evidence" / digest
        text_path = self.directory / "evidence" / (text_digest + ".utf8")
        _atomic_blob(body_path, body)
        _atomic_blob(text_path, decoded)
        rel_body = str(body_path.relative_to(self.directory))
        rel_text = str(text_path.relative_to(self.directory))
        try:
            self.db.execute("SAVEPOINT save_resource")
            existing = self.db.execute(
                "SELECT r.id FROM resources r JOIN text_views t ON t.resource_id=r.id WHERE run_id=? AND site=? AND url=? AND kind=? AND status=? AND evidence_path=? AND collector=? AND parent_resource_id=? AND t.path=?",
                (run, site, url, kind, status, rel_body, collector, parent_resource_id, rel_text),
            ).fetchone()
            if existing:
                self.db.execute("RELEASE save_resource")
                return existing[0]
            resource = self.db.execute(
                "INSERT INTO resources(run_id,site,url,kind,status,evidence_path,error,collector,encoding,parent_resource_id) VALUES(?,?,?,?,?,?,?,?,?,?)",
                (
                    run,
                    site,
                    url,
                    kind,
                    status,
                    rel_body,
                    error,
                    collector,
                    encoding,
                    parent_resource_id,
                ),
            ).lastrowid
            self.db.execute(
                "INSERT OR IGNORE INTO text_views(resource_id,path) VALUES(?,?)",
                (resource, rel_text),
            )
            for m in mentions:
                self.db.execute(
                    "INSERT OR IGNORE INTO findings VALUES(?,?,?,?,?,?,?,?,?,?,?)",
                    (
                        resource,
                        m.host,
                        m.registered,
                        m.host_type,
                        m.start,
                        m.end,
                        m.line,
                        m.column,
                        m.source,
                        m.method,
                        m.evidence,
                    ),
                )
            self.db.execute("RELEASE save_resource")
            self.db.commit()
            return resource
        except BaseException:
            self.db.execute("ROLLBACK TO save_resource")
            self.db.execute("RELEASE save_resource")
            self.db.rollback()
            raise

    @staticmethod
    def _safe_cell(value, excel=False):
        if not isinstance(value, str):
            return value
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
        if excel and len(value) > 32767:
            value = value[:32700] + "…[truncated]"
        if value[:1] in ("=", "+", "-", "@", "\t", "\r"):
            value = "'" + value
        return value

    def export(self, run: int | None = None):
        if self.db.in_transaction:
            raise ValueError("finish or rollback collection before exporting")
        if (
            run is not None
            and not self.db.execute("SELECT 1 FROM runs WHERE id=?", (run,)).fetchone()
        ):
            raise ValueError("export run does not exist")
        generation = uuid.uuid4().hex
        root = self.directory / "generations"
        root.mkdir(exist_ok=True)
        staging = root / (".tmp-" + generation)
        staging.mkdir()
        queries = {
            "域名汇总": "SELECT r.run_id,r.site,f.host,f.registered,f.host_type,COUNT(*) AS occurrences FROM findings f JOIN resources r ON r.id=f.resource_id WHERE f.host_type IN ('dns','ip','single_label') GROUP BY r.run_id,r.site,f.host,f.host_type",
            "发现明细": "SELECT r.run_id,r.id AS resource_id,r.collector,r.status,r.encoding,r.error,r.site,r.url,f.host,f.host_type,f.line,f.col,f.source,f.method,f.evidence,r.evidence_path,t.path AS decoded_text_path,f.start,f.end FROM findings f JOIN resources r ON r.id=f.resource_id LEFT JOIN text_views t ON t.resource_id=r.id ORDER BY r.run_id,r.site,r.url,f.start",
            "扫描覆盖": "SELECT q.run_id,q.site,q.state,COUNT(*) AS urls,GROUP_CONCAT(DISTINCT q.error) AS errors FROM queue q GROUP BY q.run_id,q.site,q.state",
            "引用关系": "SELECT * FROM edges ORDER BY run_id,site,parent,child",
        }
        wb = Workbook(write_only=True)
        counts = {}
        try:
            self.db.execute("BEGIN")
            for title, query in queries.items():
                if run is not None:
                    alias = (
                        "r."
                        if title in ("域名汇总", "发现明细")
                        else "q."
                        if title == "扫描覆盖"
                        else ""
                    )
                    marker = " GROUP BY" if " GROUP BY" in query else " ORDER BY"
                    clause = (
                        f" AND {alias}run_id=?" if " WHERE " in query else f" WHERE {alias}run_id=?"
                    )
                    query = query.replace(marker, clause + marker)
                cursor = self.db.execute(query, (run,) if run is not None else ())
                headers = [d[0] for d in cursor.description]
                sheet = wb.create_sheet(title)
                sheet.freeze_panes = "A2"
                sheet.append(headers)
                row_count, part = 1, 1
                counts[title] = 0
                csv_path = staging / f"{title}.csv"
                with csv_path.open("w", encoding="utf-8-sig", newline="") as stream:
                    writer = csv.writer(stream)
                    writer.writerow(headers)
                    for record in cursor:
                        if row_count >= EXCEL_ROWS:
                            part += 1
                            sheet = wb.create_sheet(f"{title}_{part}")
                            sheet.freeze_panes = "A2"
                            sheet.append(headers)
                            row_count = 1
                        values = list(record)
                        writer.writerow([self._safe_cell(v) for v in values])
                        cells = []
                        for value in values:
                            cell = WriteOnlyCell(sheet, self._safe_cell(value, excel=True))
                            if isinstance(value, str):
                                cell.data_type = "s"
                            cells.append(cell)
                        sheet.append(cells)
                        row_count += 1
                        counts[title] += 1
            self.db.rollback()
            wb.save(staging / "results.xlsx")
            target = root / generation
            os.replace(staging, target)
            manifest = self.directory / "current.json"
            temp_manifest = self.directory / (".current-" + generation)
            temp_manifest.write_text(
                json.dumps(
                    {
                        "generation": generation,
                        "path": f"generations/{generation}",
                        "run_id": run,
                        "rows": counts,
                        "complete": True,
                        "csv": "spreadsheet-safe; leading formula characters escaped",
                    }
                ),
                encoding="utf-8",
            )
            os.replace(temp_manifest, manifest)
            return target
        except BaseException:
            self.db.rollback()
            for sheet in wb.worksheets:
                if not sheet.closed:
                    sheet.close()
            shutil.rmtree(staging, ignore_errors=True)
            raise

    def close(self):
        self.db.close()
