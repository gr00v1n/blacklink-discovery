import json
import threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

from openpyxl import load_workbook

from blacklink_discovery.collect import crawl
from blacklink_discovery.extract import extract
from blacklink_discovery.store import Store


def test_mentions_and_original_locations():
    text = """// 老平台：https://legacy.example.com
// old-api.example.com
const api = "https://api.example.com/v1";
const host = "cdn." + "example.com";
/* https://partner.example.com */
const escaped = "escaped\\u002eexample.com";
"""
    mentions = extract(text, "js")
    assert {
        "legacy.example.com",
        "old-api.example.com",
        "api.example.com",
        "cdn.example.com",
        "partner.example.com",
        "escaped.example.com",
    } <= {m.host for m in mentions}
    legacy = next(m for m in mentions if m.host == "legacy.example.com")
    assert legacy.source == "JS_COMMENT" and legacy.line == 1
    concat = next(m for m in mentions if m.host == "cdn.example.com")
    assert text[concat.start : concat.end] == '"cdn." + "example.com"'
    assert concat.line == 4
    assert any(m.host_type == "single_label" for m in extract("https://intranet/help", "text"))
    assert any(m.host_type == "ip" for m in extract("http://[::1]/", "text"))
    assert any(m.host == "xn--fsqu00a.xn--0zwm56d" for m in extract("例子.测试", "text"))
    inline = extract(
        '<html>\n<script>\n// inline.example.net\nconst h="cdn."+"example.net";\n</script></html>',
        "html",
    )
    assert any(
        m.host == "inline.example.net" and m.source == "JS_COMMENT" and m.line == 3 for m in inline
    )
    assert any(m.host == "cdn.example.net" and m.line == 4 for m in inline)
    assert len(extract("https://one.example.net/path", "text")) == 1
    assert len(extract("one.example.net one.example.net", "text")) == 2
    assert any(
        m.host == "legacy.example.net" for m in extract("Old host: legacy.example.net.", "text")
    )


def test_scan_resume_resources_scope_and_export(tmp_path):
    requests = []
    pages = {
        "/": (
            "text/html",
            '<!-- archive.example.net --><a href="/page">more</a><script src="/app.js"></script><form method="post" action="/submit"></form><a href="http://localhost:1/forbidden">external</a>',
        ),
        "/page": ("text/html", '<script src="/app.js"></script><p>article.example.net</p>'),
        "/app.js": (
            "application/javascript",
            '// legacy.example.net\nconst host="cdn."+"example.net";\n//# sourceMappingURL=/app.js.map',
        ),
        "/app.js.map": ("application/json", json.dumps({"sourcesContent": ["// map.example.net"]})),
        "/robots.txt": ("text/plain", ""),
        "/sitemap.xml": ("application/xml", "<urlset/>"),
    }

    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            requests.append(("GET", self.path))
            mime, body = pages.get(self.path, ("text/plain", "missing"))
            self.send_response(200 if self.path in pages else 404)
            self.send_header("Content-Type", mime + "; charset=utf-8")
            self.end_headers()
            self.wfile.write(body.encode())

        def do_POST(self):
            requests.append(("POST", self.path))
            self.send_response(500)
            self.end_headers()

        def log_message(self, *args):
            pass

    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    store = Store(str(tmp_path))
    site = f"http://127.0.0.1:{server.server_port}/"
    try:
        run = store.new_run("test")
        crawl(store, run, [site], 6, 0, set(), 100000, max_urls=1)
        assert (
            store.db.execute("SELECT COUNT(*) FROM queue WHERE state='pending'").fetchone()[0] > 0
        )
        crawl(store, run, [site], 6, 0, set(), 100000)
        hosts = {row[0] for row in store.db.execute("SELECT host FROM findings")}
        assert {
            "archive.example.net",
            "legacy.example.net",
            "cdn.example.net",
            "map.example.net",
            "article.example.net",
        } <= hosts
        assert (
            store.db.execute(
                "SELECT COUNT(*) FROM edges WHERE child LIKE '%/app.js' AND relation='script.src'"
            ).fetchone()[0]
            == 2
        )
        assert requests.count(("GET", "/app.js")) == 1
        assert all(method == "GET" for method, _ in requests)
        assert (
            store.db.execute("SELECT COUNT(*) FROM queue WHERE state='skipped_scope'").fetchone()[0]
            == 1
        )
        target = store.export()
        book = load_workbook(target / "results.xlsx", read_only=True)
        assert book.sheetnames == ["域名汇总", "发现明细", "扫描覆盖", "引用关系"]
        assert sum(1 for _ in book["发现明细"].rows) > 5
        book.close()
    finally:
        store.close()
        server.shutdown()
        server.server_close()
        thread.join()
