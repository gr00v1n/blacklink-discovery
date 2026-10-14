import json

import pytest

from blacklink_discovery.store import Store


def test_revisions_deduplicate_without_losing_findings(tmp_path):
    store = Store(str(tmp_path))
    run = store.new_run("revision")
    url = "https://example.test/app.js"
    first = store.save(run, "site", url, b"// old.example", "// old.example", "js")
    assert store.save(run, "site", url, b"// old.example", "// old.example", "js") == first
    second = store.save(run, "site", url, b"// new.example", "// new.example", "js")
    assert second != first
    assert store.db.execute("select count(*) from resources").fetchone()[0] == 2
    assert (
        store.db.execute("select count(*) from findings where resource_id=?", (first,)).fetchone()[
            0
        ]
        == 1
    )
    store.close()


def test_failed_extract_leaves_no_resource_or_findings(tmp_path, monkeypatch):
    import blacklink_discovery.store as module

    store = Store(str(tmp_path))
    run = store.new_run("rollback")
    monkeypatch.setattr(module, "extract", lambda *_: (_ for _ in ()).throw(RuntimeError("broken")))
    with pytest.raises(RuntimeError):
        store.save(run, "site", "url", b"body", "text", "js")
    assert store.db.execute("select count(*) from resources").fetchone()[0] == 0
    store.close()


def test_legacy_migration_keeps_ids_and_export_manifest(tmp_path):
    import sqlite3

    db = sqlite3.connect(tmp_path / "discovery.sqlite")
    db.executescript("""CREATE TABLE runs(id INTEGER PRIMARY KEY,started TEXT DEFAULT CURRENT_TIMESTAMP,label TEXT);
    INSERT INTO runs(label) VALUES('legacy');
    CREATE TABLE resources(id INTEGER PRIMARY KEY,run_id INTEGER,site TEXT,url TEXT,kind TEXT,status INTEGER,evidence_path TEXT,error TEXT,UNIQUE(run_id,site,url));
    INSERT INTO resources VALUES(41,1,'site','url','js',200,'evidence/a','');
    CREATE TABLE edges(run_id INTEGER,site TEXT,parent TEXT,child TEXT,relation TEXT,UNIQUE(run_id,site,parent,child,relation));
    CREATE TABLE findings(resource_id INTEGER,host TEXT,registered TEXT,host_type TEXT,start INTEGER,end INTEGER,line INTEGER,col INTEGER,source TEXT,method TEXT,evidence TEXT,UNIQUE(resource_id,host,start,end,source));
    CREATE TABLE text_views(resource_id INTEGER PRIMARY KEY,path TEXT);
    CREATE TABLE queue(run_id INTEGER,site TEXT,url TEXT,depth INTEGER,asset INTEGER,state TEXT DEFAULT 'pending',error TEXT,UNIQUE(run_id,site,url));
    CREATE TABLE run_configs(run_id INTEGER PRIMARY KEY,config TEXT);
    CREATE TABLE run_events(id INTEGER PRIMARY KEY,run_id INTEGER,time TEXT DEFAULT CURRENT_TIMESTAMP,event TEXT,detail TEXT);
    CREATE TABLE attempts(run_id INTEGER,site TEXT,url TEXT,count INTEGER DEFAULT 0,UNIQUE(run_id,site,url));""")
    db.commit()
    db.close()
    store = Store(str(tmp_path))
    assert store.db.execute("select id from resources").fetchone()[0] == 41
    store.export(run=1)
    manifest = json.loads((tmp_path / "current.json").read_text())
    assert (tmp_path / manifest["path"] / "results.xlsx").is_file()
    store.close()


def test_atomic_export_scope_formula_safety_and_split(tmp_path, monkeypatch):
    import csv

    from openpyxl import load_workbook

    import blacklink_discovery.store as module

    store = Store(str(tmp_path))
    run = store.new_run("export")
    other = store.new_run("other")
    try:
        for r in (run, other):
            text = '=HYPERLINK("https://formula.example.net", "x")\nsecond.example.net\nthird.example.net'
            store.save(r, "site", "url", text.encode(), text, "text")
            store.edge(r, "site", "url", "child", "test")
            store.db.execute(
                "INSERT INTO queue(run_id,site,url,depth,asset,state) VALUES(?,?,?,0,0,'done')",
                (r, "site", "url"),
            )
            store.db.commit()
        monkeypatch.setattr(module, "EXCEL_ROWS", 3)
        target = store.export(run)
        with (target / "发现明细.csv").open(encoding="utf-8-sig") as stream:
            rows = list(csv.DictReader(stream))
        assert {row["run_id"] for row in rows} == {str(run)}
        assert any(row["evidence"].startswith("'=HYPERLINK") for row in rows)
        for title in ("扫描覆盖", "引用关系"):
            with (target / f"{title}.csv").open(encoding="utf-8-sig") as stream:
                assert {row["run_id"] for row in csv.DictReader(stream)} == {str(run)}
        book = load_workbook(target / "results.xlsx", read_only=True)
        assert "发现明细_2" in book.sheetnames
        assert all(cell.data_type != "f" for sheet in book for row in sheet for cell in row)
        book.close()
        before = (tmp_path / "current.json").read_bytes()
        monkeypatch.setattr(
            module.Workbook,
            "save",
            lambda *_: (_ for _ in ()).throw(OSError("simulated disk failure")),
        )
        with pytest.raises(OSError):
            store.export(run)
        assert (tmp_path / "current.json").read_bytes() == before
        assert (target / "results.xlsx").is_file()
        assert not list((tmp_path / "generations").glob(".tmp-*"))
    finally:
        store.close()
